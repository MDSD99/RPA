"""Bounded regular files under a private runner root, with atomic replacement."""

import csv
import io
import json
import os
import re
import stat
import zipfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from uuid import uuid4

from app.workflows.registry import WorkflowFailure

MAX_FILE_BYTES = 1_048_576
MAX_RESULT_BYTES = 8192
MAX_ROWS = 200
MAX_COLUMNS = 32
_RESERVED = re.compile(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", re.I)


def relative_parts(value: str) -> list[str]:
    if (
        not isinstance(value, str)
        or not value
        or len(value) > 240
        or "\\" in value
        or ":" in value
        or any(ord(char) < 32 for char in value)
    ):
        raise WorkflowFailure("FLOW_FILE_PATH_BLOCKED")
    parts = value.split("/")
    if any(
        part in {"", ".", ".."} or part.endswith((" ", ".")) or _RESERVED.match(part)
        for part in parts
    ):
        raise WorkflowFailure("FLOW_FILE_PATH_BLOCKED")
    return parts


def bounded_result(value: Any) -> Any:
    if len(json.dumps(value, ensure_ascii=False).encode("utf-8")) > MAX_RESULT_BYTES:
        raise WorkflowFailure("FLOW_FILE_RESULT_TOO_LARGE")
    return value


def table_rows(value: Any) -> list[list[Any]]:
    if not isinstance(value, list) or len(value) > MAX_ROWS:
        raise WorkflowFailure("FLOW_TABLE_ROWS_INVALID")
    for row in value:
        if not isinstance(row, list) or len(row) > MAX_COLUMNS:
            raise WorkflowFailure("FLOW_TABLE_ROWS_INVALID")
        for cell in row:
            if cell is not None and type(cell) not in {str, int, float, bool}:
                raise WorkflowFailure("FLOW_TABLE_CELL_INVALID")
            if isinstance(cell, str) and cell.lstrip().startswith(("=", "+", "-", "@")):
                raise WorkflowFailure("FLOW_TABLE_FORMULA_BLOCKED")
    bounded_result(value)
    return value


class RootedFiles:
    def __init__(self, root: Path) -> None:
        self.root = root.absolute()

    @staticmethod
    def _plain(path: Path) -> None:
        info = path.lstat()
        if stat.S_ISLNK(info.st_mode) or getattr(info, "st_file_attributes", 0) & 0x400:
            raise WorkflowFailure("FLOW_FILE_PATH_BLOCKED")

    @contextmanager
    def _parent(
        self, path: str, *, create: bool = False
    ) -> Iterator[tuple[int | None, str | Path]]:
        parts = relative_parts(path)
        descriptors = []
        try:
            self.root.mkdir(mode=0o700, parents=True, exist_ok=True)
            self._plain(self.root)
            if os.name == "posix":
                flags = os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW
                descriptor = os.open(self.root, flags)
                descriptors.append(descriptor)
                for part in parts[:-1]:
                    if create:
                        try:
                            os.mkdir(part, mode=0o700, dir_fd=descriptor)
                        except FileExistsError:
                            pass
                    descriptor = os.open(part, flags, dir_fd=descriptor)
                    descriptors.append(descriptor)
                yield descriptor, parts[-1]
            else:
                # Windows rejects symlinks and junction/reparse points. The
                # workspace additionally belongs to the current user's ACL.
                parent = self.root
                for part in parts[:-1]:
                    parent /= part
                    if create:
                        parent.mkdir(mode=0o700, exist_ok=True)
                    self._plain(parent)
                target = parent / parts[-1]
                if target.exists() or target.is_symlink():
                    self._plain(target)
                if not parent.resolve().is_relative_to(self.root.resolve()):
                    raise WorkflowFailure("FLOW_FILE_PATH_BLOCKED")
                yield None, target
        except OSError:
            raise WorkflowFailure("FLOW_FILE_UNAVAILABLE") from None
        finally:
            for descriptor in reversed(descriptors):
                os.close(descriptor)

    def read(self, path: str, check_active: Callable[[], Any]) -> bytes:
        with self._parent(path) as (parent, name):
            flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_NONBLOCK", 0)
            descriptor = os.open(name, flags, dir_fd=parent)
            with os.fdopen(descriptor, "rb") as source:
                info = os.fstat(source.fileno())
                if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
                    raise WorkflowFailure("FLOW_FILE_PATH_BLOCKED")
                if info.st_size > MAX_FILE_BYTES:
                    raise WorkflowFailure("FLOW_FILE_TOO_LARGE")
                check_active()
                data = source.read(MAX_FILE_BYTES + 1)
                check_active()
                if len(data) > MAX_FILE_BYTES:
                    raise WorkflowFailure("FLOW_FILE_TOO_LARGE")
                return data

    def write(self, path: str, data: bytes, check_active: Callable[[], Any]) -> None:
        if len(data) > MAX_FILE_BYTES:
            raise WorkflowFailure("FLOW_FILE_TOO_LARGE")
        with self._parent(path, create=True) as (parent, name):
            try:
                existing = os.stat(name, dir_fd=parent, follow_symlinks=False)
                if not stat.S_ISREG(existing.st_mode) or existing.st_nlink != 1:
                    raise WorkflowFailure("FLOW_FILE_PATH_BLOCKED")
            except FileNotFoundError:
                pass
            temporary = f".multirpa-{uuid4().hex}.tmp"
            if parent is None:
                temporary = str(Path(name).parent / temporary)
            try:
                flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0)
                descriptor = os.open(temporary, flags, mode=0o600, dir_fd=parent)
                with os.fdopen(descriptor, "wb") as target:
                    check_active()
                    target.write(data)
                    target.flush()
                    os.fsync(target.fileno())
                check_active()
                os.replace(temporary, name, src_dir_fd=parent, dst_dir_fd=parent)
                if parent is not None:
                    os.fsync(parent)
                check_active()
            finally:
                try:
                    os.unlink(temporary, dir_fd=parent)
                except FileNotFoundError:
                    pass


def encode_table(kind: str, rows: Any, sheet: str | None = None) -> bytes:
    rows = table_rows(rows)
    if kind == "csv_write":
        buffer = io.StringIO(newline="")
        csv.writer(buffer).writerows(rows)
        return buffer.getvalue().encode("utf-8")
    try:
        import defusedxml  # noqa: F401
        from openpyxl import Workbook
    except ImportError:
        raise WorkflowFailure("FLOW_EXCEL_NOT_INSTALLED") from None
    workbook = Workbook(write_only=True)
    try:
        page = workbook.create_sheet(sheet)
        for row in rows:
            page.append(row)
        binary = io.BytesIO()
        workbook.save(binary)
        return binary.getvalue()
    finally:
        workbook.close()


def decode_table(
    kind: str, data: bytes, *, sheet: str | None, check_active: Callable[[], Any]
) -> list[list[Any]]:
    if kind == "csv_read":
        try:
            source = io.StringIO(data.decode("utf-8-sig"), newline="")
            rows: list[list[Any]] = []
            for row in csv.reader(source):
                check_active()
                if len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise WorkflowFailure("FLOW_TABLE_ROWS_INVALID")
                rows.append(row)
            return bounded_result(rows)
        except (UnicodeError, csv.Error):
            raise WorkflowFailure("FLOW_CSV_INVALID") from None
    try:
        import defusedxml  # noqa: F401
        from openpyxl import load_workbook
        from openpyxl.xml import DEFUSEDXML

        if not DEFUSEDXML:
            raise WorkflowFailure("FLOW_EXCEL_XML_PROTECTION_REQUIRED")
    except ImportError:
        raise WorkflowFailure("FLOW_EXCEL_NOT_INSTALLED") from None
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 64 or sum(member.file_size for member in members) > 2_097_152:
                raise WorkflowFailure("FLOW_EXCEL_ARCHIVE_TOO_LARGE")
            if any(
                "vbaProject" in member.filename or "externalLinks/" in member.filename
                for member in members
            ):
                raise WorkflowFailure("FLOW_EXCEL_ACTIVE_CONTENT_BLOCKED")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            if sheet not in workbook.sheetnames:
                raise WorkflowFailure("FLOW_EXCEL_SHEET_MISSING")
            page = workbook[sheet]
            if (page.max_row or 0) > MAX_ROWS or (page.max_column or 0) > MAX_COLUMNS:
                raise WorkflowFailure("FLOW_TABLE_ROWS_INVALID")
            rows = []
            for excel_row in page.iter_rows():
                check_active()
                if len(rows) >= MAX_ROWS or len(excel_row) > MAX_COLUMNS:
                    raise WorkflowFailure("FLOW_TABLE_ROWS_INVALID")
                values = []
                for cell in excel_row:
                    value = cell.value
                    if value is not None and type(value) not in {str, int, float, bool}:
                        value = value.isoformat() if hasattr(value, "isoformat") else str(value)
                    values.append(value)
                rows.append(values)
            return bounded_result(rows)
        finally:
            workbook.close()
    except WorkflowFailure:
        raise
    except Exception:
        raise WorkflowFailure("FLOW_EXCEL_INVALID") from None
